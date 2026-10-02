"""Export a paper Markdown file to DOCX in the manuscript's Methods in Ecology and Evolution layout.

Arial 11 pt and double spacing on A4 portrait pages, without line numbers. Each table after
the main text goes on its own landscape page with its title above and its notes below; the
supporting information puts every table on landscape pages. Figures are 6.25 inches wide.

    python scripts/paper_documents.py docs/paper/manuscript.md manuscript.docx
    python scripts/paper_documents.py docs/paper/supporting-information.md si.docx --all-landscape

It needs pandoc and python-docx.
"""

import argparse
import copy
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Emu, Inches, Pt

A4 = (Emu(7560310), Emu(10692130))


def page(section, landscape: bool) -> None:
    width, height = A4
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (height, width) if landscape else (width, height)
    margin = Inches(0.79) if landscape else Inches(1)
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin


def set_style(style, size: float, spacing: float, bold=None, italic=None) -> None:
    style.font.name = "Arial"
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(key), "Arial")
    style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    style.font.color.rgb = None
    pf = style.paragraph_format
    pf.line_spacing = spacing
    pf.space_after = Pt(6)
    pf.space_before = Pt(0)


def table_borders(table) -> None:
    tbl = table._tbl
    pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    pr.append(borders)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--all-landscape", action="store_true")
    args = parser.parse_args()
    with tempfile.TemporaryDirectory() as tmp:
        raw = Path(tmp) / "raw.docx"
        subprocess.run(
            [
                "pandoc",
                str(args.source),
                "-f",
                "markdown-auto_identifiers",
                "-t",
                "docx",
                "-o",
                str(raw),
                f"--resource-path={args.source.parent}",
            ],
            check=True,
        )
        document = Document(raw)
    by_name = {s.name: s for s in document.styles}
    for name, size, spacing, bold, italic in [
        ("Normal", 11, 2.0, None, None),
        ("Body Text", 11, 2.0, None, None),
        ("First Paragraph", 11, 2.0, None, None),
        ("Block Text", 11, 2.0, None, None),
        ("Compact", 8, 1.0, None, None),
        ("Heading 1", 16, 1.0, True, None),
        ("Heading 2", 13, 1.0, True, None),
        ("Heading 3", 11, 1.0, True, None),
        ("Image Caption", 11, 2.0, None, False),
        ("Captioned Figure", 11, 1.0, None, None),
    ]:
        if name in by_name:
            set_style(by_name[name], size, spacing, bold, italic)

    body = document.element.body
    for table in document.tables:
        table_borders(table)
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        # Keep "< 0.0001" on one line in narrow columns.
                        r.text = r.text.replace("< ", "<\u00a0")
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8)
                        r.font.name = "Arial"
    # Table notes and table titles: single spacing, 9 pt, so a table and its note share a page.
    for p in document.paragraphs:
        text = p.text.strip()
        if text.startswith("Notes:") or (text.startswith("Table ") and p.runs and p.runs[0].bold):
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.size = Pt(9)
    # Tables before the first numbered table title stay in the portrait main text (Box 1).
    portrait_tables = set()
    for element in body.iterchildren():
        if element.tag == qn("w:tbl"):
            portrait_tables.add(element)
        elif element.tag == qn("w:p"):
            text = "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()
            if text.startswith("Table ") and (text[6:7].isdigit() or text.startswith("Table S")):
                break
    # Column widths in proportion to their content, over the text width.
    for table in document.tables:
        columns = len(table.columns)
        weights = [7.0] * columns
        for row in table.rows:
            for i, cell in enumerate(row.cells[:columns]):
                words = cell.text.replace("<\u00a0", "<_").split()
                longest = max((len(w) for w in words), default=0)
                weights[i] = max(weights[i], min(len(cell.text), 30), longest + 2)
        total = Inches(6.2 if table._tbl in portrait_tables else 10.0)
        grid = table._tbl.tblGrid
        for i, col in enumerate(grid.findall(qn("w:gridCol"))):
            col.set(qn("w:w"), str(int(total / 635 * weights[i] / sum(weights))))
        for row in table.rows:
            for i, cell in enumerate(row.cells[:columns]):
                cell.width = int(total * weights[i] / sum(weights))
        pr = table._tbl.tblPr
        for old_w in pr.findall(qn("w:tblW")):
            pr.remove(old_w)
        width = OxmlElement("w:tblW")
        width.set(qn("w:w"), str(int(total / 635)))
        width.set(qn("w:type"), "dxa")
        pr.append(width)
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        pr.append(layout)
    for shape in document.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = Inches(6.25)
        shape.height = int(Inches(6.25) * ratio)

    # Sections: the main text is portrait; each table from the first "**Table 1.**" title
    # (or every table, with --all-landscape) gets a landscape section of its own.
    paragraphs = list(document.paragraphs)
    starts = []
    for p in paragraphs:
        text = p.text.strip()
        if (
            text.startswith("Table ")
            and p.runs
            and p.runs[0].bold
            and (args.all_landscape or not text.startswith("Table S"))
        ):
            if text[6:7].isdigit() or text.startswith("Table S"):
                starts.append(p)
    first = document.sections[0]
    page(first, landscape=args.all_landscape and not starts)
    # Insert a section break before each table title: the section ending there takes the
    # orientation of what precedes it.
    for index, title in enumerate(starts):
        previous = title._p.getprevious()
        breaker = OxmlElement("w:p")
        previous.addnext(breaker)
        ppr = OxmlElement("w:pPr")
        breaker.append(ppr)
        sect = copy.deepcopy(document.sections[-1]._sectPr)
        ppr.append(sect)
        section = Section(sect, document.part)
        page(section, landscape=index > 0)
        section.start_type = WD_SECTION.NEW_PAGE
    page(document.sections[-1], landscape=bool(starts))
    for section in document.sections:
        section.start_type = WD_SECTION.NEW_PAGE
    document.save(args.output)
    print(
        f"{args.output}: {len(document.sections)} sections, {len(document.tables)} tables, "
        f"{len(document.inline_shapes)} figures",
        file=sys.stderr,
    )


if __name__ == "__main__":
    main()
